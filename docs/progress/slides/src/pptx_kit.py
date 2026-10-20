"""進捗報告スライドを pptx で組むための部品。

座標は 1920×1080 px のキャンバスで書き、先週までと同じ 10 × 5.625 in のスライドに変換する
（1 in = 192 px、1 px = 0.375 pt）。テーマは直近の pptx を土台にして引き継ぐ。

必要なもの: python-pptx（lxml は一緒に入る）
"""
import os
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

REPO = Path(__file__).resolve().parents[4]
SLIDES_DIR = REPO / "docs/progress/slides"
FONT = "Meiryo"

INK, GRAY, RULE = "1A1A1A", "6B6B6B", "C8C8C8"
BLUE, NAVY, ORANGE, ORANGE_TXT = "1F6FB4", "1F4E79", "E07B1A", "B0560C"
PALE_BLUE, PALE_ORANGE, PALE_GRAY = "E8F0F8", "FBEBDD", "F4F4F4"

__all__ = [
    "REPO", "SLIDES_DIR", "INK", "GRAY", "RULE", "BLUE", "NAVY", "ORANGE", "ORANGE_TXT",
    "PALE_BLUE", "PALE_ORANGE", "PALE_GRAY", "MSO_ANCHOR", "PP_ALIGN",
    "Deck", "P", "textbox", "box", "line", "freeform", "title", "bullets", "takeaway", "cover",
    "MILESTONES", "roadmap",
]

# docs/milestones.md の一覧と同じもの。あちらを書き換えたらここも合わせる
MILESTONES = [
    ("M1", "計測系が回る", "11/5", "6 か所の MVC と数分の記録を、API で 1 回の来室で取りきれる"),
    ("M2", "エイム中の\n差が見える", "11/26", "エイム課題の記録をイベントの時刻と合わせ、条件のあいだの差を図で示せる"),
    ("M3", "日をまたいで\n比べられる", "12/17", "自分の複数日の記録を、%MVC で日をまたいで比べられる"),
    ("M4", "実験1 の\n設計が決まる", "12/24", "予備実験で、手順・人数・来訪回数が決まる"),
    ("M5", "実験1 が\n終わる", "2027/2", "本実験と分析が終わる"),
    ("M6", "修論の\n下書き", "2027/3", "実験1 までで、修論の章立てと図表がそろう"),
    ("M7", "修論の提出", "2028/1〜2", "M6 の下書きを仕上げて出す"),
]


def px(v):
    return Emu(int(round(v * 914400 / 192)))


def fpt(v):
    return Pt(v * 0.375)


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    """直近の pptx のテーマを引き継いだ、空のスライド一式。"""

    def __init__(self, base: Path):
        self.prs = Presentation(str(base))
        ids = self.prs.slides._sldIdLst
        for sld_id in list(ids):
            self.prs.part.drop_rel(sld_id.rId)
            ids.remove(sld_id)
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "Blank")

    def slide(self, notes: str = ""):
        s = self.prs.slides.add_slide(self.blank)
        s.notes_slide.notes_text_frame.text = notes.strip()
        return s

    def save(self, out: Path):
        self.prs.save(str(out))
        print("saved", out, len(self.prs.slides), "slides")
        export_pdf(out)


def export_pdf(pptx: Path):
    """PowerPoint で PDF に書き出し、そのフルパスをクリップボードに入れる（Windows のみ）。

    OneNote の「挿入 → 印刷イメージ」でファイル名の欄に貼り付ければ、そのまま選べる。
    PowerPoint がない環境では何もしない。
    """
    if sys.platform != "win32":
        return
    pdf = SLIDES_DIR / "pdf" / (pptx.stem + ".pdf")
    pdf.parent.mkdir(exist_ok=True)
    ps = (f"$pp = New-Object -ComObject PowerPoint.Application; "
          f"$p = $pp.Presentations.Open('{pptx.resolve()}', $true, $false, $false); "
          f"$p.SaveAs('{pdf.resolve()}', 32); $p.Close(); $pp.Quit(); "
          f"Set-Clipboard -Value '{pdf.resolve()}'")
    r = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True)
    if r.returncode != 0:
        print("PDF の書き出しに失敗:", r.stderr.strip())
        return
    print("saved", pdf, "（パスをクリップボードにコピーした）")
    start_otp_watch()


# OneNote の検証コードを 10 分だけ待ち受けるスクリプト（研究とは別に置いている。なければ何もしない）
OTP_WATCH = Path.home() / "Documents/develop/ms-otp/otp_watch.py"


def start_otp_watch():
    if not OTP_WATCH.exists():
        return
    log = open(OTP_WATCH.with_name("otp_watch.log"), "a", encoding="utf-8")
    flags = subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP
    subprocess.Popen([sys.executable, str(OTP_WATCH)], stdout=log, stderr=log, creationflags=flags,
                     env={**os.environ, "PYTHONUTF8": "1"})
    print("検証コードを 10 分だけ待ち受ける（OneNote を開いてよい）")


def _set_font(run, size, bold=False, color=INK):
    f = run.font
    f.name = FONT
    f.size = fpt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # 日本語にもメイリオを当てる
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def _fill_text(tf, paras, anchor=None, pad=(0, 0, 0, 0)):
    tf.word_wrap = True
    tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = (px(p) for p in pad)
    if anchor:
        tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.line_spacing = p.get("spacing", 1.2)
        para.space_before = fpt(p.get("before", 0))
        para.space_after = fpt(p.get("after", 0))
        for text, size, bold, color in p["runs"]:
            r = para.add_run()
            r.text = text
            _set_font(r, size, bold, color)


def P(text, size, bold=False, color=INK, **kw):
    """1 段落。kw: align, before, after（px）, spacing（行間の倍率）。"""
    return dict(runs=[(text, size, bold, color)], **kw)


def textbox(slide, x, y, w, h, paras, anchor=None):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    _fill_text(tb.text_frame, paras, anchor)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, line_w=2, radius=0, dash=False,
        paras=None, anchor=MSO_ANCHOR.TOP, pad=(0, 0, 0, 0)):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = fpt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        s.line.fill.background()
    if paras:
        _fill_text(s.text_frame, paras, anchor, pad)
    return s


def line(slide, x1, y1, x2, y2, color, w=2, dash=False, head=None):
    """直線。head="end" で (x2, y2) 側、"start" で (x1, y1) 側に矢じり。"""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    style = c._element.find(qn("p:style"))
    if style is not None:  # テーマの影を付けない
        c._element.remove(style)
    c.line.color.rgb = rgb(color)
    c.line.width = fpt(w)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        el = etree.SubElement(ln, qn("a:tailEnd" if head == "end" else "a:headEnd"))
        el.set("type", "triangle")
        el.set("w", "med")
        el.set("len", "med")
    return c


def freeform(slide, pts, ox, oy, fill, stroke, w):
    """(ox, oy) を原点とする点列を結んだ多角形。"""
    b = slide.shapes.build_freeform(px(ox + pts[0][0]), px(oy + pts[0][1]), scale=1.0)
    b.add_line_segments([(px(ox + x), px(oy + y)) for x, y in pts[1:]], close=True)
    s = b.convert_to_shape()
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(stroke)
    s.line.width = fpt(w)
    return s


def cover(slide, date_text, heading="エイムスキル学習支援の研究", sub="進捗報告"):
    textbox(slide, 144, 350, 1632, 110, [P(heading, 90, True)])
    textbox(slide, 144, 534, 1632, 60, [P(sub, 50, color=GRAY)])
    line(slide, 144, 657, 566, 657, RULE, 2)
    textbox(slide, 144, 695, 1632, 50, [P(date_text, 35, color=GRAY)])


def title(slide, text, page):
    """左上の見出しと右上のページ番号。"""
    textbox(slide, 116, 80, 1500, 100, [P(text, 80, True)])
    textbox(slide, 1670, 100, 122, 40, [P(str(page), 32, color=GRAY, align=PP_ALIGN.RIGHT)])


def bullets(slide, x, y, w, h, items, size=46, gap=20):
    textbox(slide, x, y, w, h, [P("・" + t, size, True, spacing=1.3, after=gap) for t in items])


def roadmap(slide, now, page):
    """マイルストーンを横に並べ、now 番目（0 始まり）に「今ここ」を付ける。表紙の次に置く。"""
    title(slide, "マイルストーン", page)
    y, d, lw = 440, 56, 230  # 線の高さ、丸の直径、ラベルの幅
    xs = [210 + i * 236 for i in range(6)] + [1700]
    line(slide, xs[0], y, xs[5], y, RULE, 4)
    line(slide, xs[5], y, xs[6], y, RULE, 4, dash=True)
    textbox(slide, xs[5] + 40, y + 20, xs[6] - xs[5] - 80, 50,
            [P("就活・実験2（任意）", 22, True, GRAY, align=PP_ALIGN.CENTER)])
    for i, (key, name, due, _) in enumerate(MILESTONES):
        x = xs[i]
        hot, done = i == now, i < now
        box(slide, x - d / 2, y - d / 2, d, d, fill=BLUE if hot else GRAY if done else "FFFFFF",
            line=None if hot or done else RULE, line_w=3, radius=d / 2)
        if hot:
            box(slide, x - 70, y - 150, 140, 48, fill=BLUE, radius=24, anchor=MSO_ANCHOR.MIDDLE,
                paras=[P("今ここ", 24, True, "FFFFFF", align=PP_ALIGN.CENTER)])
        color = BLUE if hot else GRAY if done else INK
        textbox(slide, x - lw / 2, y - 92, lw, 40, [P(key, 30, True, color, align=PP_ALIGN.CENTER)])
        textbox(slide, x - lw / 2, y + 46, lw, 140, [
            P(name, 28, True, color, align=PP_ALIGN.CENTER, spacing=1.15, after=6),
            P(due, 24, color=GRAY, align=PP_ALIGN.CENTER)])
    key, name, due, done_when = MILESTONES[now]
    box(slide, 116, 680, 1676, 130, fill=PALE_BLUE, radius=12, anchor=MSO_ANCHOR.MIDDLE, pad=(0, 40, 0, 40),
        paras=[P(f"{key}　{name.replace(chr(10), '')}（〜{due}）", 34, True, NAVY, after=6),
               P("完了基準：" + done_when, 28, color=INK)])


def takeaway(slide, text, color=INK):
    """下の区切り線と、太字のまとめ 1 行。"""
    line(slide, 119, 822, 1792, 822, RULE, 2)
    textbox(slide, 119, 860, 1673, 70, [P(text, 46, True, color)])
